"""Document ingestion end-to-end for PDF / DOCX / MD / TXT."""
import os
os.environ["LLM_BASE_URL"] = "http://localhost:8080"
os.environ["LLM_API_KEY"]  = "test"

from pathlib import Path
from dotenv import load_dotenv; load_dotenv()
from decisiongraph.document_handlers import read_document

# Build fixtures for the formats we don't already have
Path(".bench").mkdir(exist_ok=True)
Path(".bench/doc_test.txt").write_text(
    "DecisionGraph stores institutional memory. Decisions link to topics. "
    "Topics form communities. Communities summarise organisational knowledge.",
    encoding="utf-8")
Path(".bench/doc_test.md").write_text(
    "# Architecture Decision: SQLite for code graph\n\n"
    "## Context\nWe needed fast structural queries over 1M+ symbols.\n\n"
    "## Decision\nUse SQLite with WAL + careful indexing.\n\n"
    "## Consequences\n- Sub-millisecond queries\n- No external DB dependency\n",
    encoding="utf-8")

# Try to make a real docx with python-docx if available
try:
    from docx import Document
    d = Document()
    d.add_heading("Streaming Response Headers", 0)
    d.add_paragraph("In Flask, the streaming response requires headers to be "
                     "set before the first byte is sent. Otherwise Werkzeug "
                     "raises an error. See PR #5341 for the original fix.")
    d.save(".bench/doc_test.docx")
    docx_made = True
except ImportError:
    docx_made = False

# Test reads
samples = [".bench/doc_test.txt", ".bench/doc_test.md", ".bench/test.pdf"]
if docx_made: samples.append(".bench/doc_test.docx")

print(f"=== document_handlers.read_document on {len(samples)} formats ===")
for p in samples:
    try:
        text = read_document(p)
        text_str = text if isinstance(text, str) else " ".join(
            t if isinstance(t, str) else "" for t in (text or []))
        ext = os.path.splitext(p)[1]
        print(f"  [{ext:6}] {p}")
        print(f"           length={len(text_str)} chars  "
              f"head={text_str[:80]!r}")
    except Exception as e:
        print(f"  [{p}] FAILED: {type(e).__name__}: {e}")

# Now full ingest_document via MCP handler — proves it indexes into DG
print("\n=== full ingest path: MD file -> DG graph ===")
import asyncio
from decisiongraph.mcp_server import _h_ingest_document
from decisiongraph.core import DecisionGraph
from decisiongraph.company import EnterpriseHub
from decisiongraph.discussion import DiscussionManager
dg  = DecisionGraph(storage_dir=".bench/doc_ingest_ws")
hub = EnterpriseHub(); dm = DiscussionManager()

r = asyncio.run(_h_ingest_document(dg, hub, dm, path=".bench/doc_test.md"))
print(f"  result keys: {list(r.keys()) if isinstance(r, dict) else r}")
print(f"  ok: {r.get('ok')}  stats: {r.get('stats')}")
print(f"  nodes after: {dg.G.number_of_nodes() if dg.G is not None else 'no graph'}")

# Query it
ans = dg.query("what database technology was chosen and why?")
print(f"\n  query answer (first 200): {ans[:200]}")
